import csv
import io

from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.accounts.mixins import StaffRequiredMixin
from apps.audit.services import log_export

from . import services
from .forms import DashboardFilterForm, ReportForm

PREVIEW_ROWS = 50


class DashboardView(LoginRequiredMixin, View):
    def get(self, request):
        form = DashboardFilterForm(request.GET or None)
        start, end, program, partner = form.range()
        data = services.dashboard(
            request.user, start=start, end=end, program=program, partner=partner
        )
        chart_data = {
            k: data[k]
            for k in (
                "over_time",
                "by_program",
                "needs",
                "ages",
                "races",
                "ethnicities",
                "languages",
                "zips",
            )
        }
        chart_data["partner_completion"] = [
            (p["partner"], p["completion_pct"])
            for p in data["partners"]
            if p["completion_pct"] is not None and p["received"]
        ][:10]
        return render(
            request,
            "reports/dashboard.html",
            {
                "form": form,
                "data": data,
                "chart_data": chart_data,
                "start": start,
                "end": end,
                "active": "dashboard",
                "is_partner": request.user.role == "PARTNER_USER",
            },
        )


class ReportView(StaffRequiredMixin, View):
    def get(self, request):
        form = ReportForm(request.GET or None, initial={"columns": services.DEFAULT_FIELDS})
        start, end, program, partner = form.range()
        qs = services.scoped_referrals(
            request.user, start=start, end=end, program=program, partner=partner
        ).order_by("-submitted_at")
        status = form.cleaned_data.get("status") if form.is_valid() else ""
        if status:
            qs = qs.filter(status=status)
        fields = (
            form.cleaned_data.get("columns") if form.is_valid() else None
        ) or services.DEFAULT_FIELDS
        export = request.GET.get("export")
        if export in ("csv", "xlsx"):
            headers, rows = services.report_rows(qs, fields)
            log_export(
                request,
                entity_type="Referral",
                row_count=len(rows),
                filters={k: v for k, v in request.GET.items() if v and k != "export"},
            )
            return self._export(export, headers, rows)
        headers, rows = services.report_rows(qs[:PREVIEW_ROWS], fields)
        params = request.GET.copy()
        params.pop("export", None)
        return render(
            request,
            "reports/report.html",
            {
                "form": form,
                "headers": headers,
                "rows": rows,
                "total": qs.count(),
                "preview_limit": PREVIEW_ROWS,
                "querystring": params.urlencode(),
                "active": "reports",
                "has_identifiers": bool(set(fields) & services.IDENTIFYING_FIELDS),
            },
        )

    def _export(self, kind, headers, rows):
        stamp = timezone.localdate().isoformat()
        if kind == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(headers)
            for row in rows:  # neutralize spreadsheet formula injection
                writer.writerow([_safe_cell(v) for v in row])
            resp = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
            resp["Content-Disposition"] = f'attachment; filename="referrals-{stamp}.csv"'
            return resp
        wb = Workbook()
        ws = wb.active
        ws.title = "Referrals"
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_safe_cell(v) for v in row])
        for i, header in enumerate(headers, start=1):
            width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 50)
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        resp = HttpResponse(
            out.getvalue(),
            content_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        )
        resp["Content-Disposition"] = f'attachment; filename="referrals-{stamp}.xlsx"'
        return resp


def _safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value
