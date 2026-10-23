"""Spreadsheet import for clients (CSV or Excel). Works for real, not a mockup."""

import csv
import io
import re
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from apps.clients.models import Client
from apps.referrals.services import find_or_create_client, match_client

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000

TARGET_FIELDS = {
    "first_name": "First name *",
    "last_name": "Last name *",
    "date_of_birth": "Date of birth *",
    "phone": "Phone",
    "email": "Email",
    "zip_code": "ZIP code",
    "address_line": "Street address",
    "city": "City",
    "preferred_language": "Preferred language",
    "sex": "Sex",
    "ethnicity": "Ethnicity",
}
REQUIRED = {"first_name", "last_name", "date_of_birth"}
ALIASES = {
    "first_name": ["first", "firstname", "first name", "given name", "fname"],
    "last_name": ["last", "lastname", "last name", "surname", "family name", "lname"],
    "date_of_birth": ["dob", "birth date", "birthdate", "date of birth", "birthday"],
    "phone": ["phone", "phone number", "telephone", "cell", "mobile"],
    "email": ["email", "e-mail", "email address"],
    "zip_code": ["zip", "zipcode", "zip code", "postal code"],
    "address_line": ["address", "street", "street address", "address 1"],
    "city": ["city", "town"],
    "preferred_language": ["language", "preferred language", "primary language"],
    "sex": ["sex", "gender"],
    "ethnicity": ["ethnicity", "hispanic"],
}
DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%d-%b-%Y", "%b %d, %Y"]


class ImportError_(Exception):
    pass


def parse_upload(name: str, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    if len(content) > MAX_BYTES:
        raise ImportError_("The file is larger than 5 MB.")
    if name.lower().endswith(".xlsx") or content[:4] == b"PK\x03\x04":
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter, [])]
        raw = [list(r) for r in rows_iter]
    elif name.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        headers = [h.strip() for h in next(reader, [])]
        raw = list(reader)
    else:
        raise ImportError_("Upload a .csv or .xlsx file.")
    if not any(headers):
        raise ImportError_("The first row must contain column headings.")
    rows = []
    for r in raw:
        if not any(v not in (None, "") for v in r):
            continue
        rows.append(
            {headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers)) if headers[i]}
        )
        if len(rows) > MAX_ROWS:
            raise ImportError_(f"Imports are limited to {MAX_ROWS} rows per file.")
    return [h for h in headers if h], rows


def guess_mapping(headers: list[str]) -> dict[str, str]:
    """Return {target_field: header} using common column names."""
    norm = {re.sub(r"[^a-z0-9 ]", "", h.lower()).strip(): h for h in headers}
    mapping: dict[str, str] = {}
    for target, names in ALIASES.items():
        for n in names:
            if n in norm and norm[n] not in mapping.values():
                mapping[target] = norm[n]
                break
    return mapping


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


SEX_MAP = {"f": "F", "female": "F", "m": "M", "male": "M", "x": "X", "intersex": "X"}
ETHNICITY_MAP = {
    "hispanic": "HISPANIC",
    "hispanic or latino": "HISPANIC",
    "yes": "HISPANIC",
    "not hispanic": "NOT_HISPANIC",
    "not hispanic or latino": "NOT_HISPANIC",
    "no": "NOT_HISPANIC",
    "non-hispanic": "NOT_HISPANIC",
}


def validate(rows: list[dict[str, Any]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    """One result per row: {'row': n, 'data': {...} or None, 'errors': [...], 'status': ...}."""
    missing = REQUIRED - {k for k, v in mapping.items() if v}
    if missing:
        raise ImportError_(
            "Map these required fields: "
            + ", ".join(TARGET_FIELDS[m].rstrip(" *") for m in sorted(missing))
        )
    results = []
    today = date.today()
    for i, row in enumerate(rows, start=2):  # row 1 is the heading

        def val(field, row=row):
            header = mapping.get(field)
            v = row.get(header, "") if header else ""
            return "" if v is None else (v if isinstance(v, date | datetime) else str(v).strip())

        errors, data = [], {}
        for f in ("first_name", "last_name"):
            data[f] = val(f)[:100]
            if not data[f]:
                errors.append(f"{TARGET_FIELDS[f].rstrip(' *')} is missing")
        dob = _parse_date(val("date_of_birth"))
        if dob is None:
            errors.append("Date of birth is missing or not a date (use YYYY-MM-DD or MM/DD/YYYY)")
        elif dob > today or dob.year < 1900:
            errors.append("Date of birth is out of range")
        data["date_of_birth"] = dob
        phone = val("phone")
        if phone and not re.fullmatch(r"[0-9()+\-.\s]{7,20}", phone):
            errors.append("Phone number looks invalid")
        data["phone"] = phone
        email = val("email")
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append("Email looks invalid")
        data["email"] = email
        zip_code = val("zip_code")
        if zip_code and not re.fullmatch(r"\d{5}(-\d{4})?", zip_code.split(".")[0].zfill(5)):
            errors.append("ZIP code should be 5 digits")
        data["zip_code"] = zip_code.split(".")[0].zfill(5) if zip_code else ""
        data["address_line"] = val("address_line")[:255]
        data["city"] = val("city")[:100]
        data["preferred_language"] = (val("preferred_language") or "English")[:50]
        data["sex"] = SEX_MAP.get(val("sex").lower(), "U")
        data["ethnicity"] = ETHNICITY_MAP.get(val("ethnicity").lower(), "")
        status = "error" if errors else "new"
        if not errors:
            exact, possible = match_client(data)
            status = "existing" if exact else ("possible_duplicate" if possible else "new")
        results.append(
            {
                "row": i,
                "data": None if errors else data,
                "errors": errors,
                "status": status,
                "name": f"{data['first_name']} {data['last_name']}",
            }
        )
    return results


def run_import(results: list[dict[str, Any]], *, program, user) -> dict[str, int]:
    counts = {"created": 0, "matched": 0, "flagged": 0, "skipped": 0}
    for r in results:
        if not r["data"]:
            counts["skipped"] += 1
            continue
        before = Client.objects.count()
        client = find_or_create_client(r["data"], program=program, user=user)
        extra = {
            k: r["data"][k]
            for k in ("address_line", "city", "sex", "ethnicity")
            if r["data"].get(k)
        }
        if extra:
            for k, v in extra.items():
                setattr(client, k, v)
            client.save()
        if Client.objects.count() > before:
            counts["created"] += 1
            if client.possible_duplicate_of_id:
                counts["flagged"] += 1
        else:
            counts["matched"] += 1
    return counts
